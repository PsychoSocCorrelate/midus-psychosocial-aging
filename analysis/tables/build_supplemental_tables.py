#!/usr/bin/env python3
"""Build the paper's Supplementary Tables S1-S18 as XLSX.

One sheet per table, laid out to APA 7: the table number in bold, the title in
italic title case beneath it, column spanners over grouped columns, the body
with stub section rows, and the note at the bottom (general note first, then
specific notes keyed by superscript letters). Each table is written to its own
workbook and to `supplementary_tables_combined.xlsx`, which opens on a Contents
sheet that mirrors the SI Tables index.

A table holds only what the SI or the main text states or needs to read those
statements. Table numbers follow the manuscript supplement.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "manuscript" / "supplement"

if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))


def read_csv(path: str | Path) -> pd.DataFrame:
    return pd.read_csv(ROOT / path)


# ── Labels and configuration ────────────────────────────────────────────────
def load_config() -> dict:
    with open(ROOT / "config" / "variable_config.yaml", encoding="utf-8") as f:
        return yaml.safe_load(f)


_LABELS = yaml.safe_load((ROOT / "config" / "predictor_labels.yaml").read_text(encoding="utf-8"))
LABEL: dict[str, str] = _LABELS["labels"]
CONTRASTS: dict[str, dict] = _LABELS["contrasts"]


def label(var: str) -> str:
    """Reader-facing label of a predictor, or of one indicator contrast of a categorical one."""
    var = str(var)
    if "__" in var:
        base, level = var.split("__", 1)
        return f"{LABEL[base]}: {CONTRASTS[base]['levels'][level]}"
    return LABEL[var]


def tier_maps() -> tuple[dict[str, int], dict[str, str]]:
    cfg = load_config()
    tiers = {v: 1 for v in cfg.get("tier1", [])} | {v: 2 for v in cfg.get("tier2", [])}
    subs = {var: sd for sd, variables in cfg.get("subdomain", {}).items() for var in variables}
    return tiers, subs


SUBDOMAIN_LABEL = {
    "tier_1": "Background (Tier 1)",
    "personality": "Personality",
    "psychological_wellbeing": "Psychological well-being",
    "mental_health_symptoms": "Mental health symptoms",
    "coping_resilience": "Coping and resilience",
    "life_evaluation": "Life evaluation",
    "social_relationships": "Social relationships",
    "intimate_partner": "Intimate partnership",
    "household_caregiving": "Household and caregiving",
    "health_behaviors": "Health behaviors",
    "cognitive_activities": "Cognitive activities",
    "substance_use": "Substance use",
    "work_occupation": "Work and occupation",
    "socioeconomic_subjective": "Subjective socioeconomic resources",
    "healthcare_access": "Healthcare access",
}

# Stem of the allostatic-load outcome columns.
AL_STEM = "ALcomp"
OUTCOME_STEMS = [("FInoPsy", "Frailty"), ("btactComp", "Cognition"), (AL_STEM, "Allostatic load")]
OUTCOME_LABEL = dict(OUTCOME_STEMS)


def analyzed_predictors() -> list[str]:
    """The 177 predictors in the order the SI lists them: Tier 1, then Tier 2 by subdomain."""
    reliability = read_csv("output/d1_descriptive/predictor_reliability.csv")
    analyzed = set(reliability["Variable"])
    cfg = load_config()
    order = [v for v in cfg["tier1"] if v in analyzed]
    for sd in SUBDOMAIN_LABEL:
        if sd != "tier_1":
            order += [v for v in cfg["subdomain"][sd] if v in analyzed]
    if len(order) != 177 or set(order) != analyzed:
        raise ValueError(f"expected the 177 analyzed predictors, got {len(order)}")
    return order


def grouped_predictors() -> list[tuple[str, list[str]]]:
    """(section heading, predictors) in SI order, each heading carrying its count."""
    order = analyzed_predictors()
    tiers, subs = tier_maps()
    groups = []
    t1 = [v for v in order if tiers[v] == 1]
    groups.append((f"Tier 1: background ({len(t1)})", t1))
    for sd, lab in SUBDOMAIN_LABEL.items():
        if sd == "tier_1":
            continue
        members = [v for v in order if tiers[v] == 2 and subs.get(v) == sd]
        groups.append((f"Tier 2: {lab.lower()} ({len(members)})", members))
    return groups


# ── Number formatting (APA: no leading zero where a value cannot exceed 1) ──
_LEAD = re.compile(r"^([+-]?)0\.")
_SUP = str.maketrans("-0123456789", "⁻⁰¹²³⁴⁵⁶⁷⁸⁹")


def num(x, d: int = 3, lead: bool = True, sign: bool = False) -> str:
    if x is None or (isinstance(x, float) and np.isnan(x)) or pd.isna(x):
        return ""
    s = f"{x:+.{d}f}" if sign else f"{x:.{d}f}"
    if s.lstrip("+-").strip("0.") == "":
        s = s.lstrip("+-")  # no signed zero
    return _LEAD.sub(r"\1.", s) if not lead else s


def sig(x, n: int = 3, lead: bool = True) -> str:
    """n significant figures; powers of ten written out below .001."""
    if pd.isna(x):
        return ""
    if x == 0:
        return "0"
    if abs(x) < 1e-3:
        mant, exp = f"{x:.{n - 1}e}".split("e")
        return f"{mant} × 10{str(int(exp)).translate(_SUP)}"
    d = max(n - 1 - int(np.floor(np.log10(abs(x)))), 0)
    return num(x, d, lead=lead)


def sig_ci(x, lo, hi, n: int = 3) -> str:
    """An estimate and its interval at the decimals that give the largest of them n significant figures."""
    if pd.isna(x):
        return ""
    ref = max(abs(x), abs(lo), abs(hi))
    if ref < 1e-3:
        return f"{sig(x, n)} [{sig(lo, n)}, {sig(hi, n)}]"
    d = max(n - 1 - int(np.floor(np.log10(ref))), 0)
    return f"{num(x, d)} {interval(lo, hi, d, True)}"


def pval(p) -> str:
    """Exact p to three decimals; below .001 as a power of ten, as the text cites them."""
    if pd.isna(p):
        return ""
    if p < 1e-3:
        mant, exp = f"{p:.1e}".split("e")
        return f"{mant} × 10{str(int(exp)).translate(_SUP)}"
    return num(p, 3, lead=False)


def interval(lo, hi, d: int = 3, lead: bool = False) -> str:
    if pd.isna(lo) or pd.isna(hi):
        return ""
    return f"[{num(lo, d, lead)}, {num(hi, d, lead)}]"


def thousands(n) -> str:
    return "" if pd.isna(n) else f"{int(n):,}"


# ── The APA table model and its writer ──────────────────────────────────────
@dataclass
class Col:
    head: str
    width: float = 10
    fmt: str | None = None   # Excel number format for numeric cells
    left: bool = False       # stub and text columns align left


@dataclass
class Section:
    """A stub heading spanning the table, in italic."""
    text: str


@dataclass
class Span:
    """A cell value merged across `n` columns."""
    value: object
    n: int


@dataclass
class Block:
    cols: list[Col]
    rows: list
    spanners: list[list[tuple[str, int]]] = field(default_factory=list)
    panel: str | None = None


@dataclass
class SITable:
    key: str
    filename: str
    title: str
    blocks: list[Block]
    note: str
    specific: list[str] = field(default_factory=list)
    bold: set = field(default_factory=set)   # (block, row, col) cells set in bold
    freeze: bool = False


FONT = "Arial"
THIN = Side(style="thin", color="000000")
R3 = ".000;-.000;.000"        # bounded statistics, no leading zero
R2 = ".00;-.00;.00"
F3 = "0.000;-0.000;0.000"
INT = "#,##0"


def _rule(ws, row: int, c0: int, c1: int, top: bool = False) -> None:
    for c in range(c0, c1 + 1):
        cell = ws.cell(row=row, column=c)
        b = cell.border
        cell.border = Border(top=THIN if top else b.top, bottom=b.bottom if top else THIN,
                             left=b.left, right=b.right)


def _row_height(texts_by_width: list[tuple[str, float]], base: float = 13.0) -> float:
    lines = 1
    for text, width in texts_by_width:
        if not text:
            continue
        chars = max(width * 1.15, 4)
        n = sum(max(1, int(np.ceil(len(part) / chars))) for part in str(text).split("\n"))
        lines = max(lines, n)
    return base * lines


def display_round(value, fmt: str):
    """Store a number at the precision its format shows, rounded as Python formats it.

    The text's numbers were formatted in Python from the same outputs; Excel rounds the
    decimal value half away from zero, so .6545 would show as .655 where the text says
    .654. Storing the Python-rounded value makes the cell and the text agree.
    """
    first = fmt.split(";")[0]
    d = len(first.split(".")[1]) if "." in first else 0
    if first.startswith("#"):
        return value
    return float(f"{value:.{d}f}") if d else int(f"{value:.0f}")


def write_table(ws, t: SITable) -> None:
    ws.sheet_view.showGridLines = False
    ncols = max(len(b.cols) for b in t.blocks)
    widths = [0.0] * ncols
    for b in t.blocks:
        for j, c in enumerate(b.cols):
            widths[j] = max(widths[j], c.width)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    table_width = sum(widths)

    ws.cell(row=1, column=1, value=f"Table {t.key}").font = Font(name=FONT, bold=True, size=11)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    ws.cell(row=2, column=1, value=t.title).font = Font(name=FONT, italic=True, size=11)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    ws.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = _row_height([(t.title, table_width)], 15)
    r = 4
    header_end = None
    title_rows = None
    for bi, b in enumerate(t.blocks):
        if b.panel:
            cell = ws.cell(row=r, column=1, value=b.panel)
            cell.font = Font(name=FONT, bold=True, size=10)
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            r += 1
        first_header = r
        title_rows = title_rows or (first_header, None)
        for level in b.spanners:
            c = 1
            for text, span in level:
                if text:
                    cell = ws.cell(row=r, column=c, value=text)
                    cell.font = Font(name=FONT, size=10)
                    cell.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
                    if span > 1:
                        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
                    _rule(ws, r, c, c + span - 1)
                c += span
            r += 1
        for j, col in enumerate(b.cols, start=1):
            cell = ws.cell(row=r, column=j, value=col.head)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(horizontal="left" if col.left else "center",
                                       vertical="bottom", wrap_text=True)
        ws.row_dimensions[r].height = _row_height([(c.head, c.width) for c in b.cols])
        _rule(ws, first_header, 1, ncols, top=True)
        _rule(ws, r, 1, ncols)
        header_end = header_end or r
        r += 1
        for ri, row in enumerate(b.rows):
            if isinstance(row, Section):
                cell = ws.cell(row=r, column=1, value=row.text)
                cell.font = Font(name=FONT, italic=True, size=10)
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
                r += 1
                continue
            c = 1
            texts = []
            for value in row:
                span = 1
                if isinstance(value, Span):
                    value, span = value.value, value.n
                col = b.cols[c - 1]
                if isinstance(value, (np.generic,)):
                    value = value.item()
                if isinstance(value, float) and np.isnan(value):
                    value = None
                cell = ws.cell(row=r, column=c, value=value)
                cell.font = Font(name=FONT, size=10, bold=(bi, ri, c - 1) in t.bold)
                numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
                if numeric and col.fmt:
                    cell.number_format = col.fmt
                    cell.value = display_round(value, col.fmt)
                cell.alignment = Alignment(horizontal="left" if col.left and span == 1 else "center",
                                           vertical="top", wrap_text=True)
                if span > 1:
                    ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
                width = sum(widths[c - 1:c - 1 + span])
                texts.append(("" if value is None else str(value), width))
                c += span
            ws.row_dimensions[r].height = _row_height(texts)
            r += 1
        _rule(ws, r - 1, 1, ncols)
        if bi < len(t.blocks) - 1:
            r += 1
    # Note, at the bottom.
    notes = [CellRichText(TextBlock(InlineFont(i=True, rFont=FONT, sz=9), "Note."),
                          TextBlock(InlineFont(rFont=FONT, sz=9), " " + t.note))]
    notes += [CellRichText(TextBlock(InlineFont(rFont=FONT, sz=9), s)) for s in t.specific]
    for rich in notes:
        cell = ws.cell(row=r, column=1, value=rich)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        ws.row_dimensions[r].height = _row_height([(str(rich), table_width * 1.05)], 12)
        r += 1
    if t.freeze and header_end:
        ws.freeze_panes = f"B{header_end + 1}"
    ws.page_setup.orientation = "landscape" if table_width > 110 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    if t.freeze and header_end:
        ws.print_title_rows = f"{title_rows[0]}:{header_end}"


# ── Table S1, predictor inventory ───────────────────────────────────────────
SPARSE_FLOOR_FLAG_OVERRIDES = {"homeless", "sexualOrientation", "anhedonia", "anxietyScore",
                               "careGivIntensity"}
STATISTIC = {"Pearson r": "r", "Spearman rho": "ρ", "Cohen's kappa": "κ"}


def complete_mz_pair_values(df: pd.DataFrame, var: str) -> pd.DataFrame:
    required = {"pair_id", "sib_type", var}
    if not required.issubset(df.columns):
        return pd.DataFrame(columns=["member_1", "member_2"])
    cols = ["pair_id", "sib_type", var] + (["M2MRID"] if "M2MRID" in df.columns else [])
    mask = df["sib_type"].eq("MZ") & df["pair_id"].notna()
    if "is_full_pair" in df.columns:
        mask &= df["is_full_pair"].fillna(False).astype(bool)
    sub = df.loc[mask, cols].dropna(subset=[var]).copy()
    counts = sub.groupby("pair_id").size()
    sub = sub[sub["pair_id"].isin(counts[counts == 2].index)]
    if sub.empty:
        return pd.DataFrame(columns=["member_1", "member_2"])
    sub = sub.sort_values(["pair_id"] + (["M2MRID"] if "M2MRID" in sub.columns else []))
    return sub.groupby("pair_id", sort=True)[var].agg(member_1="first", member_2="last").reset_index(drop=True)


def sparse_floor_flag(var, vtype, mz_value, mz_agreement, mz_pairs) -> bool:
    """A near-zero MZ value that arises from little variation, not from absent resemblance."""
    if var in SPARSE_FLOOR_FLAG_OVERRIDES:
        return True
    if pd.isna(mz_value):
        return False
    if vtype in {"binary", "nominal"}:
        return mz_value <= 0.05 and pd.notna(mz_agreement) and mz_agreement >= 95
    if vtype == "continuous" and not mz_pairs.empty:
        values = pd.concat([mz_pairs["member_1"], mz_pairs["member_2"]]).dropna()
        return mz_value <= 0.05 and float(values.astype(str).value_counts(normalize=True).iloc[0]) >= 0.80
    return False


def build_s1() -> SITable:
    rel = read_csv("output/d1_descriptive/predictor_reliability.csv").set_index("Variable")
    fam = read_csv("output/d1_descriptive/predictor_family_resemblance_w1.csv")
    fam = fam[fam["Sibling stratum"].eq("MZ")].set_index("Variable")
    master = pd.read_csv(ROOT / "data" / "processed" / "master_analysis.csv", low_memory=False)
    rows = []
    for heading, members in grouped_predictors():
        rows.append(Section(heading))
        for var in members:
            r = rel.loc[var]
            mz = fam.loc[var] if var in fam.index else pd.Series(dtype=object)
            mz_value = pd.to_numeric(mz.get("Value"), errors="coerce")
            flagged = sparse_floor_flag(var, r["Type"], mz_value,
                                        pd.to_numeric(mz.get("Pct_Agreement"), errors="coerce"),
                                        complete_mz_pair_values(master, var))
            mz_cell = num(mz_value, 3, lead=False) + ("ᵃ" if flagged else "") if pd.notna(mz_value) else ""
            rows.append([
                LABEL[var], var, r["Type"].capitalize(), int(r["N"]), STATISTIC[r["Metric"]],
                float(r["Value"]),
                float(r["Pct_Agreement"]) if pd.notna(r["Pct_Agreement"]) else None,
                int(mz["N_pairs"]) if pd.notna(mz.get("N_pairs")) else None,
                mz_cell,
            ])
    cfg = load_config()
    analyzed = set(rel.index)
    excluded = [(v, t) for t, key in ((1, "tier1"), (2, "tier2")) for v in cfg[key] if v not in analyzed]
    excluded_text = " and ".join(f"{LABEL[v].lower()} (Tier {t})" for v, t in excluded)
    cols = [Col("Predictor", 34, left=True), Col("Code", 22, left=True), Col("Type", 11),
            Col("n", 8, INT), Col("Statistic", 9), Col("Test-retest", 10, R3),
            Col("Agreement (%)", 11, "0.0"), Col("MZ pairs", 8, INT), Col("MZ co-twin resemblance", 13)]
    return SITable(
        "S1", "table_s1_psychosocial_exposome.xlsx",
        "Predictor Inventory: The 16 Tier 1 and 161 Tier 2 Predictors",
        [Block(cols, rows)],
        note=(
            "Predictors are grouped by tier and subdomain, with the number of predictors in "
            "parentheses; code is the variable name used in the analysis code and figures. "
            "Test-retest reliability is the Wave 1 to Wave 2 association among the n participants "
            "observed at both waves: Pearson r for continuous, Spearman ρ for ordinal, and Cohen's κ "
            "for binary and nominal predictors; agreement is the percentage of participants giving "
            "the same response at both waves. MZ co-twin resemblance is the same statistic computed "
            "between monozygotic co-twins at Wave 1. Summaries of test-retest reliability by "
            "measurement type and subdomain are in Table S12. Two configured variables, "
            f"{excluded_text}, were excluded before modeling for high missingness and are not "
            "listed. MZ = monozygotic."
        ),
        specific=["ᵃ Near zero because the predictor varies little among co-twins (at least 95% "
                  "agreement, or at least 80% of values in one category); not evidence of absent "
                  "resemblance."],
        freeze=True,
    )


# ── Table S2, frailty construct overlap ─────────────────────────────────────
def frailty_items() -> pd.DataFrame:
    return read_csv("config/frailty_index_items.csv")


def build_s2() -> SITable:
    adj = read_csv("output/a2_frailty_overlap/a2_adjusted_r2.csv").set_index("model")
    model_rows = [
        ("Tier 1 only", "T1_only"),
        ("Tier 1 + Tier 2, all predictors", "T1T2_full"),
        ("Excluding direct overlap (a)", "T1T2_excl_a"),
        ("Excluding direct and proximal overlap (a, b)", "T1T2_excl_ab"),
    ]
    panel_a = [[lab, int(adj.loc[k, "n_features"]), float(adj.loc[k, "r2_mean"]),
                float(adj.loc[k, "r2_std"]), None if k == "T1_only" else float(adj.loc[k, "delta_r2"])]
               for lab, k in model_rows]
    audit = read_csv("config/frailty_overlap_audit.csv")
    tiers, subs = tier_maps()
    analyzed = set(analyzed_predictors())
    audit = audit[audit["predictor"].isin(analyzed) & audit["overlap_code"].isin(["a", "b"])].copy()
    items = frailty_items().set_index("code")

    def item_list(codes: str) -> str:
        return "; ".join(f"{items.loc[c.strip(), 'item'].lower()} ({items.loc[c.strip(), 'number']})"
                         for c in str(codes).split(","))

    order = {v: i for i, v in enumerate(analyzed_predictors())}
    audit["_o"] = audit["predictor"].map(order)
    audit = audit.sort_values(["overlap_code", "_o"])
    panel_b = []
    for code, sub in audit.groupby("overlap_code", sort=True):
        panel_b.append(Section("Direct overlap (a)" if code == "a" else "Proximal overlap (b)"))
        for _, r in sub.iterrows():
            rationale = str(r["rationale"]).replace(" (FInoPsy item)", "")
            panel_b.append([LABEL[r["predictor"]], r["predictor"],
                            SUBDOMAIN_LABEL[subs[r["predictor"]]], item_list(r["frailty_items_involved"]),
                            rationale[0].upper() + rationale[1:]])
    cs1 = read_csv("output/cs1_cross_sectional/cs1_r2_table.csv")
    n_frailty = int(cs1[(cs1["outcome"] == "FInoPsy_W1") & (cs1["tier_config"] == "T1T2")]["n"].iloc[0])
    n_a = int((audit["overlap_code"] == "a").sum())
    n_b = int((audit["overlap_code"] == "b").sum())
    a_cols = [Col("Model", 34, left=True), Col("Predictors", 22, INT), Col("R²", 22, R3),
              Col("SD", 26, R3), Col("ΔR²", 40, R3)]
    b_cols = [Col("Predictor", 34, left=True), Col("Code", 22, left=True), Col("Subdomain", 22, left=True),
              Col("Frailty-index items (Table S14 number)", 26, left=True), Col("Rationale", 40, left=True)]
    return SITable(
        "S2", "table_s2_adjusted_frailty_results.xlsx",
        "Frailty Construct-Overlap Sensitivity Analysis",
        [Block(a_cols, panel_a, panel="Panel A. Cross-validated prediction of Wave 1 frailty"),
         Block(b_cols, panel_b, panel=f"Panel B. Tier 2 predictors coded for content overlap with "
                                     f"frailty-index items ({n_a + n_b})")],
        note=(
            "Panel A gives the cross-validated R² of the cross-sectional frailty model (n = "
            f"{n_frailty:,}) with Tier 1 only, with all 177 predictors, and after "
            f"excluding the {n_a} Tier 2 predictors coded as direct overlap and the {n_a + n_b} coded as "
            f"direct or proximal overlap (leaving {161 - n_a - n_b} Tier 2 predictors). SD is across "
            "the five folds; ΔR² is relative to Tier 1 only. Panel B lists the coded predictors. "
            "Direct overlap (a), the predictor duplicates or closely matches a frailty-index item; "
            "proximal overlap (b), the predictor is not an index item but could reflect the same "
            "self-reported health limitation or illness perception. The remaining Tier 2 predictors "
            "were coded as having no direct content overlap. ADL = activities of daily living."
        ),
    )


# ── Table S3, cross-domain transfer ─────────────────────────────────────────
PRIMARY_W1 = [f"{s}_W1" for s, _ in OUTCOME_STEMS]
PAIR_ORDER = [(a, b) for a in PRIMARY_W1 for b in PRIMARY_W1 if a != b]


def short(code: str) -> str:
    return OUTCOME_LABEL[code.replace("_W1", "")]


def build_s3() -> SITable:
    tuned = read_csv("output/a3_domain_divergence/a3_cross_prediction_cs1.csv")
    dec = read_csv("output/a3_domain_divergence/a3_cross_prediction_decomposed_cs1.csv")
    null = read_csv("output/a3_domain_divergence/a3_partition_null_cs1.csv")
    rows = []
    for src, tgt in PAIR_ORDER:
        t = tuned[(tuned["source_outcome"] == src) & (tuned["target_outcome"] == tgt)].iloc[0]
        d = dec[(dec["source_outcome"] == src) & (dec["target_outcome"] == tgt)].set_index("feature_set")
        n = null[(null["source_outcome"] == src) & (null["target_outcome"] == tgt)].iloc[0]
        shared = sorted(d.loc["shared", "features"].split(","), key=lambda v: LABEL[v])
        if int(n["n_shared"]) != len(shared):
            raise ValueError(f"{src}->{tgt}: partition null and decomposition disagree on the shared set")
        rows.append([
            short(src), short(tgt), float(t["target_own_r2"]), float(t["cross_r2"]),
            float(t["r2_ratio"]) * 100,
            int(d.loc["shared", "n_psych_features"]), float(d.loc["shared", "cross_r2"]),
            int(d.loc["unique_source", "n_psych_features"]), float(d.loc["unique_source", "cross_r2"]),
            int(d.loc["unique_target", "n_psych_features"]), float(d.loc["unique_target", "cross_r2"]),
            float(n["actual_shared_r2"]),
            f"{num(n['null_mean_r2'], 3, False)} ({num(n['null_sd_r2'], 3, False)})",
            pval(n["p_value"]),
            "; ".join(LABEL[v] for v in shared),
        ])
    lo = min(r[4] for r in rows)
    hi = max(r[4] for r in rows)
    cols = [Col("Source", 13, left=True), Col("Target", 13, left=True), Col("Target R²", 9, R3),
            Col("R²", 8, R3), Col("% of target R²", 9, "0"),
            Col("n", 5, INT), Col("R²", 7, R3), Col("n", 5, INT), Col("R²", 7, R3),
            Col("n", 5, INT), Col("R²", 7, R3),
            Col("Shared-set R²", 9, R3), Col("Random-subset R², M (SD)", 12), Col("P", 7),
            Col("Shared predictors", 46, left=True)]
    spanners = [[("", 3), ("Source top 20", 2), ("Decomposition of the top-20 sets", 6),
                 ("Random-partition null", 3), ("", 1)],
                [("", 5), ("Shared", 2), ("Source-unique", 2), ("Target-unique", 2), ("", 4)]]
    return SITable(
        "S3", "table_s3_cross_domain_prediction.xlsx",
        "Cross-Domain Transfer of Top Predictors, With Shared and Unique Decomposition and a "
        "Random-Partition Null",
        [Block(cols, rows, spanners)],
        note=(
            "Each source outcome's 20 highest-ranked predictors by bootstrap mean |SHAP| in its "
            "cross-sectional model (for frailty, the model excluding the 29 content-overlapping "
            "predictors; Table S2) were refit as 5-fold cross-validated models of the target outcome, "
            "with outcomes residualized on age within fold. Target R², the target's own full-model R² "
            "(177 predictors). Source top 20, all 20 source predictors with hyperparameters tuned as "
            f"in the primary models; their R² recovered {lo:.0f}% to {hi:.0f}% of the target's. "
            "Decomposition: shared, the source predictors also among the target's top 20; "
            "source-unique, the remaining source predictors; target-unique, the target's top 20 not "
            "among the source's, predicting the target itself; these sets were fit with fixed "
            "hyperparameters (tree depth 3, learning rate .01, 500 trees). Random-partition null: the "
            "shared set and 1,000 random subsets of the source's top 20 of the same size, each fit with "
            "the fixed hyperparameters on one 80/20 split of the target sample; P is the proportion of "
            "random subsets whose R² reached the shared set's. Each outcome's ranking is in Table S17."
        ),
    )


# ── Within-family tables, S4-S10 ────────────────────────────────────────────
MDE_FILES = {"cs": "output/within_family_mde/within_family_mde_concurrent.csv",
             "change": "output/within_family_mde/within_family_mde.csv"}
GROUP_SAMPLE = {"All_paired": "All", "MZ_sensitivity": "MZ"}
DOMAIN_CODE = {lab: stem for stem, lab in OUTCOME_STEMS}
# Household income is in dollars; its raw coefficients print per $10,000.
RAW_SCALE = {"householdIncome": 1e4}


def family_results() -> pd.DataFrame:
    from analysis.tables.build_table1 import build_full_results
    full = build_full_results()
    parts = []
    for layer, path in MDE_FILES.items():
        m = read_csv(path)[["sample", "outcome_code", "focal_var", "MDE_std"]]
        parts.append(m.assign(analysis=layer))
    mde = pd.concat(parts)
    full["sample"] = full["group"].map(GROUP_SAMPLE)
    full = full.merge(mde, how="left", left_on=["analysis", "sample", "outcome", "focal_var"],
                      right_on=["analysis", "sample", "outcome_code", "focal_var"])
    if full["MDE_std"].isna().any():
        missing = full.loc[full["MDE_std"].isna(), ["analysis", "group", "outcome", "focal_var"]]
        raise KeyError(f"no MDE for {missing.to_dict('records')}")
    return full


def icc_table() -> pd.DataFrame:
    return read_csv("output/a6_sibling/a6_icc_denominators.csv")


LAYER_TEXT = {
    "cs": dict(target="Wave 1 outcome", adjust="age (partitioned the same way) and sex",
               model="cross-sectional"),
    "change": dict(target="Wave 2 outcome",
                   adjust="the Wave 1 outcome and age (each partitioned the same way) and sex",
                   model="9-year change"),
}
CONTRAST_NOTE = (
    "Categorical predictors enter as indicator contrasts, estimated when at least two pairs were "
    "discordant: work status, retired or not working vs. employed; occupational group, "
    "professional or managerial or blue-collar vs. service, sales, or office; work industry, goods "
    "production, trade and transport, or knowledge and government vs. services and arts."
)


def build_family_table(key: str, layer: str, domain: str, full: pd.DataFrame) -> SITable:
    sub = full[(full["analysis"] == layer) & (full["domain"] == domain)]
    rows, bold = [], set()
    for group, heading in (("All_paired", "All pairs"), ("MZ_sensitivity", "MZ pairs")):
        g = sub[sub["group"] == group]
        rows.append(Section(f"{heading} ({len(g)} contrasts tested)"))
        for _, r in g.iterrows():
            base = r["focal_var"].split("__")[0]
            scale = RAW_SCALE.get(base, 1.0)
            lab = label(r["focal_var"]) + (" (per $10,000)" if scale != 1 else "")
            if r["beta_W_q"] <= 0.05:
                bold.add((0, len(rows), 6))
            b, lo, hi = (r[c] * scale for c in ("beta_W", "beta_W_ci_lo", "beta_W_ci_hi"))
            rows.append([
                lab, int(r["n_pairs"]),
                int(r["n_MZ"]) if group == "All_paired" and pd.notna(r["n_MZ"]) else None,
                sig_ci(b, lo, hi),
                float(r["beta_W_std"]), pval(r["beta_W_p"]), pval(r["beta_W_q"]), float(r["MDE_std"]),
                sig(r["beta_B"] * scale, 3), float(r["beta_B_std"]), pval(r["beta_B_p"]),
            ])
    cols = [Col("Predictor", 38, left=True), Col("Pairs", 7, INT), Col("MZ pairs", 7, INT),
            Col("βW [95% CI]", 26), Col("βW std", 8, R3),
            Col("p", 11), Col("q", 11), Col("MDE", 7, R3),
            Col("βB", 10), Col("βB std", 8, R3), Col("p", 11)]
    spanners = [[("", 3), ("Within family", 5), ("Between family", 3)]]
    text = LAYER_TEXT[layer]
    is_cs = layer == "cs"
    name = domain.title()
    title = (f"Cross-Sectional Within- and Between-Family Associations With {name}" if is_cs else
             f"Within- and Between-Family Associations With 9-Year Change in {name}")
    return SITable(
        key, FAMILY_FILES[key], title, [Block(cols, rows, spanners)],
        note=(
            f"Random-intercept mixed models of the {text['target']} on each focal predictor's pair "
            "mean (between family, βB) and deviation from the pair mean (within family, βW), adjusting "
            f"for {text['adjust']}, with a random intercept for pair. Focal predictors had bootstrap "
            f"top-20 stability of at least .60 in the {text['model']} predictive model. All pairs pools "
            "monozygotic (MZ), same-sex dizygotic, and opposite-sex dizygotic twins; pairs is the number "
            "of complete pairs in each predictor's model, and MZ pairs the number of MZ pairs among them "
            "(outcome-level pair counts are in Table S4). βW is in outcome units per predictor unit, "
            "with a 95% CI from 1,000 pair-level bootstrap resamples; βW std and βB std scale the "
            "coefficient by the within-pair (between-pair) SDs of predictor and outcome. p is from the "
            "mixed model; q is the Benjamini-Hochberg value within outcome, pair sample, and layer, in "
            "bold when q ≤ .05. MDE is the minimum detectable βW std at 80% power and two-sided "
            f"α = .05 (2.80 standard errors). {CONTRAST_NOTE}"
        ),
        bold=bold,
    )


FAMILY_FILES = {
    "S5": "table_s5_family_cs_frailty.xlsx", "S6": "table_s6_family_cs_cognition.xlsx",
    "S7": "table_s7_family_cs_allostatic_load.xlsx", "S8": "table_s8_family_change_frailty.xlsx",
    "S9": "table_s9_family_change_cognition.xlsx", "S10": "table_s10_family_change_allostatic_load.xlsx",
}


def build_s4(full: pd.DataFrame) -> SITable:
    icc = icc_table()
    rows = []
    for stem, domain in OUTCOME_STEMS:
        rows.append(Section(domain))
        for group, lab in (("All_paired", "All pairs"), ("MZ", "MZ pairs"), ("DZ_SS", "Same-sex DZ pairs")):
            row = [lab]
            survivors = {}
            for layer in ("cs", "change"):
                ic = icc[(icc["analysis"] == layer) & (icc["outcome"] == stem) & (icc["group"] == group)].iloc[0]
                model_group = {"All_paired": "All_paired", "MZ": "MZ_sensitivity"}.get(group)
                if model_group:
                    g = full[(full["analysis"] == layer) & (full["outcome"] == stem) & (full["group"] == model_group)]
                    s = g[g["beta_W_q"] <= 0.05].sort_values("beta_W_std", key=abs, ascending=False)
                    survivors[layer] = "; ".join(f"{label(v)} ({num(b, 2, False, sign=False)})"
                                                 for v, b in zip(s["focal_var"], s["beta_W_std"]))
                    cells = [f"{g['n_pairs'].min()}-{g['n_pairs'].max()}",
                             float(ic["icc"]), f"{len(s)}/{len(g)}"]
                    if layer == "cs":
                        cells.append(float(s["beta_W_std"].abs().median()) if len(s) else None)
                    cells.append(float(g["MDE_std"].median()))
                else:
                    survivors[layer] = ""
                    cells = ["", float(ic["icc"]), ""] + ([None] if layer == "cs" else []) + [None]
                row += [int(ic["n_pairs"])] + cells
            row += [survivors["cs"], survivors["change"]]
            rows.append(row)
    cols = [Col("Pair sample", 19, left=True),
            Col("Pairs", 6, INT), Col("Model pairs", 9), Col("ICC", 6, R3), Col("q ≤ .05 / tested", 8),
            Col("Median |βW std|, q ≤ .05", 9, R2), Col("Median MDE", 7, R2),
            Col("Pairs", 6, INT), Col("Model pairs", 9), Col("ICC", 6, R3), Col("q ≤ .05 / tested", 8),
            Col("Median MDE", 7, R2),
            Col("Cross-sectional", 50, left=True), Col("9-year change", 30, left=True)]
    spanners = [[("", 1), ("Cross-sectional (Wave 1)", 6), ("9-year change", 5),
                 ("Within-pair associations with q ≤ .05 (βW std)", 2)]]
    return SITable(
        "S4", "table_s4_family_cs_overview_icc.xlsx",
        "Within-Family Overview: Twin-Pair Counts, Intraclass Correlations by Zygosity, Power, and "
        "Within-Pair Associations Surviving FDR Correction",
        [Block(cols, rows, spanners)],
        note=(
            "Pairs, twin pairs with the outcome observed (cross-sectional, at Wave 1; change, at both "
            "waves); model pairs, the range of complete pairs across the focal predictors' models "
            "(Tables S5-S10). All pairs pools monozygotic (MZ), same-sex dizygotic (DZ), and "
            "opposite-sex DZ twins; no non-twin siblings are included, and DZ pairs were not modeled "
            "separately. ICC, the outcome intraclass correlation σ²u / (σ²u + σ²e) from a "
            "random-intercept model; for change, of the Wave 2 outcome conditional on Wave 1. "
            "q ≤ .05 / tested, the within-pair associations surviving Benjamini-Hochberg correction "
            "over the contrasts tested, within each outcome, pair sample, and layer; the surviving "
            "associations are listed with their βW std, by absolute size. MDE, the minimum detectable "
            "βW std at 80% power and two-sided α = .05, as a median over the contrasts tested. "
            "FDR = false discovery rate."
        ),
    )


# ── Table S11, sample ──────────────────────────────────────────────────────
def build_s11() -> SITable:
    desc = read_csv("output/d1_descriptive/sample_descriptives.csv")
    stat = dict(zip(desc["Variable"].astype(str).str.strip(), desc["Statistic"]))
    att = read_csv("output/d1_descriptive/attrition.csv").set_index("Outcome Domain")
    cs1 = read_csv("output/cs1_cross_sectional/cs1_r2_table.csv")
    a4 = read_csv("output/a4_longitudinal/a4_r2_table.csv")
    a4b = read_csv("output/a4b_longitudinal/a4b_r2_table.csv")
    a9 = read_csv("output/a9_concurrent/a9_r2_table.csv")
    icc = icc_table()

    def split(s: str) -> tuple[str, str]:
        value, _, n = str(s).partition(", N=")
        return value, n

    total = int(str(stat["N (total)"]).replace(",", ""))
    race = {"1.0": "White", "2.0": "Black", "3.0": "Asian, Pacific Islander, or Native American",
            "4.0": "Other"}
    race_n = sum(int(str(stat[k]).split(" ")[0].replace(",", "")) for k in race)
    sample_rows = [
        Section(f"Baseline sample, MIDUS 2 and MIDUS Refresher pooled (N = {total:,})"),
        ["MIDUS 2, n (%)", Span(stat["MIDUS2"], 3)],
        ["MIDUS Refresher, n (%)", Span(stat["MIDUSR"], 3)],
        ["Age, M (SD)", Span(split(stat["Age (W1)"])[0], 3)],
        ["Female, n (%)", Span(split(stat["Female"])[0], 3)],
        ["Hispanic, n (%)ᵃ", Span(split(stat["Hispanic"])[0], 3)],
    ]
    for k, lab in race.items():
        sample_rows.append([f"Race: {lab}, n (%)", Span(stat[k], 3)])
    sample_rows.append(["Race: not reported, n (%)",
                        Span(f"{total - race_n:,} ({(total - race_n) / total:.1%})", 3)])
    sample_rows.append(["Education credential (1-12), M (SD)ᵃ", Span(split(stat["Education (years)"])[0], 3)])

    att_key = {"Frailty": "Frailty Index", "Cognition": "Cognitive Composite", "Allostatic load": "Allostatic Load"}
    desc_key = {"Frailty": "Frailty Index", "Cognition": "Cognitive Composite", "Allostatic load": "Allostatic Load"}
    out_rows = [Section("Outcomes")]
    for wave in ("W1", "W2"):
        cells_m, cells_n = [], []
        for _, lab in OUTCOME_STEMS:
            m, n = split(stat[f"{desc_key[lab]} ({wave})"])
            cells_m.append(m)
            cells_n.append(int(n.replace(",", "")))
        w = "Wave 1" if wave == "W1" else "Wave 2"
        out_rows.append([f"{w}, M (SD)"] + cells_m)
        out_rows.append([f"{w}, n observed"] + cells_n)
    out_rows.append(["Observed at both waves, n (% of Wave 1)"] +
                    [f"{int(att.loc[att_key[lab], 'N_W2']):,} ({att.loc[att_key[lab], 'Retention%']:.1f}%)"
                     for _, lab in OUTCOME_STEMS])
    out_rows.append(Section("Completers vs. non-completers at baseline, d (p)"))
    for var, lab in (("Age", "Age"), ("Female", "Female"), ("Education", "Education credential")):
        out_rows.append([lab] + [f"{num(att.loc[att_key[o], f'{var}_d'], 3, sign=True)} "
                                 f"({pval(att.loc[att_key[o], f'{var}_p'])})" for _, o in OUTCOME_STEMS])
    out_rows.append(["Wave 1 outcome"] + [
        f"{num(att.loc[att_key[o], f'{att_key[o]}_W1_d'], 3, sign=True)} "
        f"({pval(att.loc[att_key[o], f'{att_key[o]}_W1_p'])})" for _, o in OUTCOME_STEMS])

    def n_of(df, outcome, tier, **where):
        sub = df[(df["outcome"] == outcome) & (df["tier_config"] == tier)]
        for k, v in where.items():
            sub = sub[sub[k] == v]
        return int(sub["n"].iloc[0])

    def pairs(group):
        cells = []
        for stem, _ in OUTCOME_STEMS:
            ns = [int(icc[(icc["analysis"] == layer) & (icc["outcome"] == stem) & (icc["group"] == group)]
                      ["n_pairs"].iloc[0]) for layer in ("change", "cs")]
            cells.append(f"{ns[0]}-{ns[1]}")
        return cells

    ana_rows = [
        Section("Analytic samples"),
        ["Cross-sectional models (Wave 1)"] + [n_of(cs1, f"{s}_W1", "T1T2") for s, _ in OUTCOME_STEMS],
        ["9-year change, Wave 1 predictors"] + [n_of(a4, f"{s}_W2", "Yw1T1T2", cohort="pooled") for s, _ in OUTCOME_STEMS],
        ["9-year change, Wave 2 predictors"] + [n_of(a4b, f"{s}_W2", "Yw1T1T2", cohort="pooled") for s, _ in OUTCOME_STEMS],
        ["Wave 2 status, Wave 2 predictors"] + [n_of(a9, f"{s}_W2", "T1T2") for s, _ in OUTCOME_STEMS],
        ["Within family, all twin pairs (pairs)"] + pairs("All_paired"),
        ["Within family, MZ pairs (pairs)"] + pairs("MZ"),
    ]
    pheno = n_of(cs1, "phenoAge_res", "T1T2")
    cols = [Col("Characteristic", 52, left=True), Col("Frailty", 18, INT), Col("Cognition", 18, INT),
            Col("Allostatic load", 18, INT)]
    return SITable(
        "S11", "table_s11_sample_descriptives_attrition.xlsx",
        "Sample Characteristics, Longitudinal Retention, and Analytic Samples",
        [Block(cols, sample_rows + out_rows + ana_rows)],
        note=(
            "Characteristics of the pooled baseline sample are n (%) or M (SD); race percentages are "
            "of the full baseline sample. Outcomes are on their analysis scales: frailty, the "
            "proportion of eligible deficits; cognition, the BTACT composite; allostatic load, the "
            "mean of 16 biomarker z scores. Completers are Wave 1 participants with the outcome "
            "observed again at Wave 2; d is the standardized baseline difference between completers "
            "and non-completers (completers minus non-completers, pooled SD), with Welch t-test p. "
            "Machine-learning samples exclude complete twin pairs for frailty and cognition and retain "
            "them for allostatic load (Table S18 reports models holding them out). Within-family "
            "ranges run from the 9-year change to the cross-sectional pair count (Table S4). The "
            f"PhenoAge benchmark sample was n = {pheno:,}. BTACT = Brief Test of Adult Cognition by "
            "Telephone; MZ = monozygotic."
        ),
        specific=[f"ᵃ Observed for {split(stat['Hispanic'])[1]} (Hispanic ethnicity) and "
                  f"{split(stat['Education (years)'])[1]} (education) participants; education is the "
                  "MIDUS 12-category credential code."],
    )


# ── Table S12, reliability ─────────────────────────────────────────────────
def build_s12() -> SITable:
    tr = read_csv("output/d1_descriptive/outcome_test_retest.csv").set_index("Outcome")
    fam = read_csv("output/d1_descriptive/outcome_family_resemblance_w1.csv")
    key = {"Frailty": "Frailty Index", "Cognition": "Cognitive Composite", "Allostatic load": "Allostatic Load"}
    rows_a = []
    for _, lab in OUTCOME_STEMS:
        r = tr.loc[key[lab]]
        f = fam[fam["Outcome"] == key[lab]].set_index("Sibling stratum")
        twins = [f"{num(f.loc[g, 'Value'], 3, False)} ({int(f.loc[g, 'N_pairs'])})"
                 for g in ("MZ", "DZ_SS", "DZ_OS", "All_paired")]
        rows_a.append([lab, int(r["N"]), float(r["r_Pearson"]), float(r["r_Spearman"]),
                       float(r["r_age_residualized"]), float(r["sd_change"])] + twins)
    cols_a = [Col("Outcome", 26, left=True), Col("n", 8, INT), Col("r", 8, R3), Col("ρ", 8, R3),
              Col("r, age-residualized", 11, R3), Col("SD of change", 9, F3),
              Col("MZ", 12), Col("Same-sex DZ", 12), Col("Opposite-sex DZ", 12), Col("All pairs", 12)]
    spanners_a = [[("", 1), ("Wave 1 to Wave 2 test-retest", 5),
                   ("Wave 1 co-twin correlation, r (pairs)", 4)]]

    rel = read_csv("output/d1_descriptive/predictor_reliability.csv")
    tiers, subs = tier_maps()
    rel["subdomain"] = rel["Variable"].map(lambda v: "tier_1" if tiers.get(v) == 1 else subs.get(v))
    rows_b = [Section("By measurement type")]
    types = [("Continuous", ["continuous"]), ("Ordinal", ["ordinal"]), ("Binary", ["binary"]),
             ("Nominal", ["nominal"]), ("Binary and nominal", ["binary", "nominal"])]
    for lab, ts in types:
        g = rel[rel["Type"].isin(ts)]
        stat = STATISTIC[g["Metric"].iloc[0]]
        agree = g["Pct_Agreement"].median()
        rows_b.append([lab, len(g), stat, float(g["Value"].median()),
                       f"{num(g['Value'].min(), 3, False)} to {num(g['Value'].max(), 3, False)}",
                       None if pd.isna(agree) else round(float(agree), 1)])
    rows_b.append(Section("By subdomain"))
    # Subdomain medians come from d1_descriptive, which produces the medians the SI reports.
    by_domain = read_csv("output/d1_descriptive/predictor_reliability_by_domain.csv").set_index("domain")
    for sd, lab in SUBDOMAIN_LABEL.items():
        g = rel[rel["subdomain"] == sd]
        if int(by_domain.loc[sd, "n_predictors"]) != len(g):
            raise ValueError(f"{sd}: d1 counts {by_domain.loc[sd, 'n_predictors']} predictors, Table S1 {len(g)}")
        mix = ", ".join(f"{STATISTIC[m]} ({n})" for m, n in g["Metric"].value_counts().items())
        rows_b.append([lab, len(g), mix, float(by_domain.loc[sd, "median_reliability"]),
                       f"{num(g['Value'].min(), 3, False)} to {num(g['Value'].max(), 3, False)}", None])
    cols_b = [Col("Predictors", 26, left=True), Col("n", 8, INT), Col("Statistic (n)", 16, left=True),
              Col("Median", 11, R3), Col("Range", 16), Col("Median agreement (%)", 12, "0.0")]
    return SITable(
        "S12", "table_s12_reliability_correlations.xlsx",
        "Test-Retest Reliability and Co-Twin Resemblance of the Outcomes, and Test-Retest "
        "Reliability of the Predictors",
        [Block(cols_a, rows_a, spanners_a, panel="Panel A. Primary outcomes"),
         Block(cols_b, rows_b, panel="Panel B. Predictor test-retest reliability, Wave 1 to Wave 2")],
        note=(
            "Panel A: test-retest correlations among participants observed at both waves, raw "
            "(Pearson r and Spearman ρ) and after residualizing each wave on age; SD of change is the "
            "standard deviation of the Wave 2 minus Wave 1 score, on the outcome's analysis scale. "
            "Co-twin correlations are Pearson r between co-twins at Wave 1, with the number of pairs "
            "in parentheses; model-based intraclass correlations are in Table S4. Panel B summarizes "
            "the 177 predictors' test-retest statistics (Table S1): Pearson r for continuous, "
            "Spearman ρ for ordinal, and Cohen's κ for binary and nominal predictors, with median "
            "percent agreement for the latter. Subdomain medians pool these statistics, with the "
            "number of predictors on each in parentheses. MZ = monozygotic; DZ = dizygotic."
        ),
    )


# ── Table S13, allostatic-load biomarkers ──────────────────────────────────
AL_SYSTEM_LABEL = {"cardiovascular": "Cardiovascular", "metabolic": "Metabolic",
                   "hpa_axis": "Hypothalamic-pituitary-adrenal axis", "inflammatory": "Inflammatory",
                   "renal": "Renal"}


def build_s13() -> SITable:
    from config.biomarker_constants import BIOMARKERS

    inv = read_csv("config/al_biomarker_inventory.csv").set_index("code")
    markers = [b[0] for b in BIOMARKERS]
    missing = sorted(set(markers) - set(inv.index))
    if missing:
        raise KeyError(f"composite markers absent from the inventory: {missing}")
    system = {b[0]: b[2] for b in BIOMARKERS}
    rows = []
    for sd, lab in AL_SYSTEM_LABEL.items():
        members = [m for m in markers if system[m] == sd]
        rows.append(Section(f"{lab} ({len(members)})"))
        for m in members:
            r = inv.loc[m]
            rows.append([r["biomarker"], r["midus2"], r["midus3"], r["refresher1"], r["risk_direction"],
                         "" if pd.isna(r["medication_adjustment"]) else r["medication_adjustment"]])
    if len(markers) != 16 or len(AL_SYSTEM_LABEL) != len(set(system.values())):
        raise ValueError("expected 16 markers in 5 systems")
    cols = [Col("Biomarker (unit)", 40, left=True), Col("MIDUS 2", 12), Col("MIDUS 3", 12),
            Col("Refresher 1", 12), Col("Risk tail", 9), Col("Medication adjustment", 40, left=True)]
    spanners = [[("", 1), ("MIDUS Biomarker Project variable", 3), ("", 2)]]
    return SITable(
        "S13", "table_s13_AL_biomarkers.xlsx",
        "Allostatic-Load Biomarkers: 16 Markers in Five Physiological Systems",
        [Block(cols, rows, spanners)],
        note=(
            "Markers are grouped by physiological system, with the number of markers in parentheses. "
            "Variables are the MIDUS Biomarker Project codes at MIDUS 2 (B4), MIDUS 3 (C4), and MIDUS "
            "Refresher 1 (RA4). Risk tail, the direction scored as dysregulated. Medication "
            "adjustments are applied before scoring for participants taking the listed medication. "
            "Each marker is standardized against its cohort's baseline mean and SD and oriented so "
            "that higher values indicate greater risk; the composite is the mean of these z scores, "
            "computed when at least 13 of the 16 are observed. BMI and waist-hip ratio are predictors, "
            "not components. HDL = high-density lipoprotein."
        ),
    )


# ── Table S14, frailty-index items ─────────────────────────────────────────
def build_s14() -> SITable:
    items = frailty_items()
    rows = []
    for _, r in items.iterrows():
        vars_ = ["" if pd.isna(r[c]) else str(r[c]).replace("/", ", ")
                 for c in ("midus2", "midus3", "refresher1", "refresher2")]
        rows.append([int(r["number"]), r["item"], r["domain"], *vars_, r["scoring"]])
    n_r2 = int(items["refresher2"].notna().sum())
    lonely = items[items["code"] == "lonely"].iloc[0]
    absent = items[items["refresher2"].isna()]["item"].str.lower().tolist()
    cols = [Col("No.", 5, INT), Col("Item", 30, left=True), Col("Domain", 18, left=True),
            Col("MIDUS 2", 16), Col("MIDUS 3", 16), Col("Refresher 1", 16), Col("Refresher 2", 16),
            Col("Deficit score", 48, left=True)]
    spanners = [[("", 3), ("MIDUS variable", 4), ("", 1)]]
    return SITable(
        "S14", "table_s14_FInoPsy_items.xlsx", "Frailty-Index Items",
        [Block(cols, rows, spanners)],
        note=(
            "The index is the proportion of eligible deficits present, each item scored from 0 (no "
            "deficit) to 1, excluding items flagged as psychological symptoms. All "
            f"{len(items)} items were collected at baseline (MIDUS 2, Refresher 1) and at the MIDUS 3 "
            f"follow-up; the Refresher 2 follow-up carried neither {' nor '.join(absent)}, so its "
            f"index is built from {n_r2} items. Where an item combines several variables, the score "
            f"column states the rule. Item {int(lonely['number'])}, loneliness, is the one psychosocially "
            "adjacent item retained (SI Methods); Table S2 lists the predictors coded for overlap with "
            "the items."
        ),
        freeze=True,
    )


# ── Table S15, outcome correlations ────────────────────────────────────────
def build_s15() -> SITable:
    df = read_csv("output/outcome_correlations_by_wave/outcome_correlations_by_wave.csv")
    name = {"Frailty Index": "Frailty", "Cognitive Composite": "Cognition", "Allostatic Load": "Allostatic load"}
    pairs = [("Frailty Index", "Cognitive Composite"), ("Frailty Index", "Allostatic Load"),
             ("Cognitive Composite", "Allostatic Load")]
    rows = []
    for a, b in pairs:
        row = [f"{name[a]} with {name[b].lower()}"]
        for wave in ("W1", "W2"):
            r = df[(df["wave"] == wave) & (df["outcome_1"] == a) & (df["outcome_2"] == b)].iloc[0]
            row += [float(r["r"]), interval(r["ci_lo"], r["ci_hi"], 2, False), int(r["n"])]
        rows.append(row)
    cols = [Col("Outcomes", 30, left=True), Col("r", 8, R2), Col("95% CI", 14), Col("n", 8, INT),
            Col("r", 8, R2), Col("95% CI", 14), Col("n", 8, INT)]
    spanners = [[("", 1), ("Baseline (Wave 1)", 3), ("9-year follow-up (Wave 2)", 3)]]
    return SITable(
        "S15", "table_s15_outcome_correlations.xlsx",
        "Age-Residualized Correlations Among the Primary Outcomes, by Wave",
        [Block(cols, rows, spanners)],
        note=("Pearson correlations between outcomes residualized on age within wave, among "
              "participants observed on both outcomes, with Fisher-z 95% confidence intervals."),
    )


# ── Table S16, profile comparison ──────────────────────────────────────────
PROFILE_SOURCE = {"P90": "90th percentile", "P10": "10th percentile", "median": "median",
                  "modal category": "modal category", "min/max mean-SHAP level": "extreme mean SHAP",
                  "override": "clinical band"}


def fmt_setting(v) -> str:
    try:
        x = float(v)
    except (TypeError, ValueError):
        return str(v)
    return f"{x:.0f}" if x == int(x) else f"{x:.2f}"


def build_s16() -> SITable:
    pred = read_csv("output/profile_comparison/profile_predictions.csv")
    shap = read_csv("output/profile_comparison/profile_shap.csv")
    spec = yaml.safe_load((ROOT / "config" / "profile_specification.yaml").read_text())
    rows = []
    for outcome in dict.fromkeys(pred["outcome_key"]):
        p = pred[pred["outcome_key"] == outcome].set_index("profile")
        name = p["outcome"].iloc[0]
        rows.append(Section(name))
        focal = {fe["name"]: fe for fe in spec["outcomes"][outcome]["focal"]}
        sub = shap[shap["outcome_key"] == outcome]
        for f in dict.fromkeys(sub["feature"]):
            by = sub[sub["feature"] == f].set_index("profile")["focal_value"]
            fe = focal.get(f, {})
            src = [(fe.get(lvl) or {}).get("source", "") for lvl in ("favorable", "neutral", "unfavorable")]
            rule = (f"{PROFILE_SOURCE.get(src[0], src[0])} / {PROFILE_SOURCE.get(src[1], src[1])} / "
                    f"{PROFILE_SOURCE.get(src[2], src[2])}")
            rows.append([LABEL[f], str(fe.get("type", "")).capitalize(), rule,
                         fmt_setting(by["favorable"]), fmt_setting(by["neutral"]),
                         fmt_setting(by["unfavorable"]), ""])
        cell = lambda k: (f"{num(p.loc[k, 'prediction'], 3, True, True)} "
                          f"{interval(p.loc[k, 'ci_lo'], p.loc[k, 'ci_hi'], 3, True)}")
        rows.append([f"Predicted {name.lower()} [95% CI]", "", "", cell("favorable"), cell("neutral"),
                     cell("unfavorable"), cell("favorable_minus_unfavorable")])
        rows.append(["Sample percentile", "", "", *[f"{p.loc[k, 'percentile']:.0f}"
                                                    for k in ("favorable", "neutral", "unfavorable")], ""])
    cols = [Col("Focal predictor", 32, left=True), Col("Type", 11, left=True),
            Col("Setting rule (favorable / neutral / unfavorable)", 44, left=True),
            Col("Favorable", 22), Col("Neutral", 22), Col("Unfavorable", 22),
            Col("Favorable − unfavorable", 22)]
    return SITable(
        "S16", "table_s16_profile_comparison.xlsx",
        "Predicted Aging Outcomes Under Favorable, Neutral, and Unfavorable Focal-Predictor Profiles",
        [Block(cols, rows)],
        note=(
            "For each outcome, the six focal predictors were set to their favorable, neutral, or "
            "unfavorable values and all other predictors held at the sample reference. Continuous "
            "predictors: neutral, the sample median; favorable and unfavorable, the 10th and 90th "
            "percentiles, oriented by the direction of the SHAP dependence (BMI uses clinical bands). "
            "Categorical predictors: neutral, the modal category; favorable and unfavorable, the "
            "categories with the most and least healthy mean SHAP contribution among categories with "
            "at least 30 observations. Predictions are from the cross-sectional model on the "
            "age-residualized scale, where zero is the sample average for a given age; 95% CIs are "
            "percentiles of 200 bootstrap refits with hyperparameters fixed, and omit model-selection "
            "variance. Sample percentile, the position of the prediction in the observed outcome "
            "distribution. Exploratory and model-implied, not causal; the allostatic-load drinking "
            "setting follows the model's direction and is not a recommendation."
        ),
    )


# ── Table S17, importance concentration and stability ─────────────────────
SHAP_SUMMARIES = {
    "Cross-sectional": ("output/cs1_cross_sectional/shap", "{stem}_W1_T1T2_bootstrap_shap.csv"),
    "9-year change": ("output/a4_longitudinal/shap", "{stem}_W2_long_bootstrap_shap.csv"),
}


def read_shap_summary(layer: str, stem: str) -> pd.DataFrame:
    folder, pattern = SHAP_SUMMARIES[layer]
    df = read_csv(Path(folder) / pattern.format(stem=stem))
    df = df.rename(columns={df.columns[0]: "Variable"})
    # The change models carry the Wave 1 outcome as a predictor; it is not one of the 177.
    df = df[df["Variable"] != f"{stem}_W1"].reset_index(drop=True)
    df = df.sort_values("mean_abs_shap", ascending=False).reset_index(drop=True)
    df["rank"] = np.arange(1, len(df) + 1)
    return df.set_index("Variable")


def crossing_rank(shap: pd.Series, share: float) -> int:
    """How many predictors, taken in descending order, first reach `share` of the total."""
    s = shap.sort_values(ascending=False)
    return int((s.cumsum() / s.sum() < share).sum() + 1)


def build_s17() -> SITable:
    tiers, _ = tier_maps()
    rho = read_csv("output/a4_longitudinal/a4_cs_vs_long_rho.csv").set_index("outcome")
    rec = read_csv("output/recurring_predictors/recurring_predictors.csv").set_index("feature")["n_lists"]
    models = [(stem, lab, layer) for stem, lab in OUTCOME_STEMS for layer in SHAP_SUMMARIES]
    shap = {(stem, layer): read_shap_summary(layer, stem) for stem, _, layer in models}

    summary = []
    def srow(text, fn):
        cells = [text]
        for stem, _, layer in models:
            v = fn(stem, layer)
            cells.append(Span("" if v is None else v, 3))
        cells.append("")
        summary.append(cells)

    t2 = lambda stem, layer: shap[(stem, layer)][shap[(stem, layer)].index.map(tiers.get) == 2]["mean_abs_shap"]
    srow("Tier 2 predictors to reach 50% of summed mean |SHAP|", lambda s, l: crossing_rank(t2(s, l), .5))
    srow("Tier 2 predictors to reach 80% of summed mean |SHAP|", lambda s, l: crossing_rank(t2(s, l), .8))
    srow("Predictors in the top 20 of at least 1 of 200 fits",
         lambda s, l: int((shap[(s, l)]["top20_stability"] > 0).sum()))
    srow("Predictors in the top 20 of at least 60% of fits",
         lambda s, l: int((shap[(s, l)]["top20_stability"] >= .6).sum()))
    srow("Predictors in the top 20 of all fits",
         lambda s, l: int((shap[(s, l)]["top20_stability"] >= 1).sum()))
    srow("Spearman ρ with the cross-sectional ranking (177 predictors)",
         lambda s, l: num(rho.loc[f"{s}_W2", "spearman_rho"], 2, False) if l != "Cross-sectional" else None)
    srow("Top-20 overlap with the cross-sectional ranking (of 19)",
         lambda s, l: int(str(rho.loc[f"{s}_W2", "top20_overlap"]).split("/")[0]) if l != "Cross-sectional" else None)

    rows = [Section("Concentration, stability, and cross-layer correspondence")] + summary
    for heading, members in grouped_predictors():
        rows.append(Section(heading))
        for var in members:
            row = [LABEL[var]]
            for stem, _, layer in models:
                r = shap[(stem, layer)].loc[var]
                row += [int(r["rank"]),
                        f"{num(r['mean_abs_shap'], 4)} {interval(r['ci_lower'], r['ci_upper'], 4, True)}",
                        float(r["top20_stability"])]
            row.append(int(rec.get(var, 0)))
            rows.append(row)
    cols = [Col("Predictor", 34, left=True)]
    for _ in models:
        cols += [Col("Rank", 5, INT), Col("Mean |SHAP| [95% CI]", 24), Col("Top-20 stability", 8, R3)]
    cols.append(Col("Top-20 lists (of 6)", 8, INT))
    spanners = [[("", 1)] + [(lab, 6) for _, lab in OUTCOME_STEMS] + [("", 1)],
                [("", 1)] + [(layer, 3) for _, _, layer in models] + [("", 1)]]
    return SITable(
        "S17", "table_s17_importance_stability.xlsx",
        "Concentration, Bootstrap Stability, and Cross-Layer Correspondence of Predictor Importance",
        [Block(cols, rows, spanners)],
        note=(
            "Cross-sectional models predict the Wave 1 outcome from the 177 predictors; 9-year change "
            "models predict the Wave 2 outcome from the Wave 1 outcome and the same 177 Wave 1 "
            "predictors, and the Wave 1 outcome, in the top 20 of every fit, is not ranked or counted. "
            "Mean |SHAP| is the mean absolute out-of-bag SHAP value on the outcome's scale over 200 "
            "bootstrap fits, with a 95% percentile interval, to four decimals; rank orders it within each model; top-20 "
            "stability is the share of fits placing the predictor in the top 20. The first rows count, "
            "among the 161 Tier 2 predictors ranked by mean |SHAP|, those needed to reach one half and "
            "four fifths of their summed mean |SHAP|, and, among all 177, those entering the top 20 in "
            "at least one, at least 60%, and all fits; Spearman ρ correlates the two layers' mean |SHAP| "
            "and top-20 overlap is counted over the 19 places left beside the Wave 1 outcome. Top-20 "
            "lists, the number of the six outcome-by-layer top-20 lists the predictor enters; the ten "
            "entering at least four are the recurring predictors (SI Appendix, Cross-Domain Shared and "
            "Unique Predictors)."
        ),
        freeze=True,
    )


# ── Table S18, robustness of the increments ────────────────────────────────
def build_s18() -> SITable:
    cs1 = read_csv("output/cs1_cross_sectional/cs1_r2_table.csv")
    a4 = read_csv("output/a4_longitudinal/a4_r2_table.csv")
    a4b = read_csv("output/a4b_longitudinal/a4b_r2_table.csv")
    a9 = read_csv("output/a9_concurrent/a9_r2_table.csv")
    a12 = read_csv("output/a12_delta_r2_ci/a12_delta_r2_ci.csv")
    tab = read_csv("output/cs1_cross_sectional/cs1_tabpfn_r2.csv").set_index("outcome")
    held = read_csv("output/al_twins_heldout/al_twins_heldout_summary.csv").set_index("layer")

    def r2(df, outcome, tier, **where):
        sub = df[(df["outcome"] == outcome) & (df["tier_config"] == tier)]
        for k, v in where.items():
            sub = sub[sub[k] == v]
        return sub.iloc[0]

    def ci(arm, outcome):
        sub = a12[(a12["arm"] == arm) & (a12["outcome"] == outcome)]
        return "" if sub.empty else interval(sub["p2.5"].iloc[0], sub["p97.5"].iloc[0])

    def spec_rows(df, pattern, y_only, base, full, arm, where, names=OUTCOME_STEMS, tabpfn=False):
        out = []
        for stem, lab in names:
            o = pattern.format(stem=stem)
            b, f = r2(df, o, base, **where), r2(df, o, full, **where)
            y = r2(df, o, y_only, **where)["r2_mean"] if y_only else None
            out.append([lab, int(f["n"]), y, float(b["r2_mean"]), float(f["r2_mean"]),
                        float(f["r2_mean"] - b["r2_mean"]), ci(arm, o) if arm else "",
                        None if y is None else float(b["r2_mean"] - y),
                        float(tab.loc[o, "r2_mean"]) if tabpfn else None])
        return out

    rows = [Section("Wave 1 status, Wave 1 predictors")]
    rows += spec_rows(cs1, "{stem}_W1", None, "T1", "T1T2", "concurrent", {}, tabpfn=True)
    rows += spec_rows(cs1, "{stem}", None, "T1", "T1T2", None, {}, names=[("phenoAge_res", "PhenoAge")],
                      tabpfn=True)
    rows.append(Section("9-year change, Wave 1 predictors"))
    rows += spec_rows(a4, "{stem}_W2", "Yw1", "Yw1T1", "Yw1T1T2", "conditional", {"cohort": "pooled"})
    rows.append(Section("9-year change, Wave 1 predictors, by cohort"))
    for stem, lab in OUTCOME_STEMS[:2]:
        for cohort, clab in (("MIDUS2", "MIDUS 2"), ("MIDUSR", "MIDUS Refresher")):
            rows += spec_rows(a4, "{stem}_W2", "Yw1", "Yw1T1", "Yw1T1T2", None, {"cohort": cohort},
                              names=[(stem, f"{lab}, {clab}")])
    rows.append(Section("9-year change, Wave 2 predictors"))
    rows += spec_rows(a4b, "{stem}_W2", "Yw1", "Yw1T1", "Yw1T1T2", "conditional_followup", {"cohort": "pooled"})
    rows.append(Section("Wave 2 status, Wave 2 predictors"))
    rows += spec_rows(a9, "{stem}_W2", None, "T1", "T1T2", None, {})
    rows.append(Section("Allostatic load with complete twin pairs held out of training"))
    for layer, lab in (("cross_sectional", "Wave 1 status"), ("longitudinal", "9-year change")):
        h = held.loc[layer]
        rows.append([lab, int(h["n_heldout"]), None, float(h["r2_base_heldout"]), float(h["r2_full_heldout"]),
                     float(h["delta_r2_T2_heldout"]), "", None, None])
    cs_held = held.loc["cross_sectional"]
    cols = [Col("Outcome", 30, left=True), Col("n", 8, INT), Col("Wave 1 outcome only", 10, R3),
            Col("Baseline model", 10, R3), Col("Full model", 10, R3), Col("ΔR² Tier 2", 9, R3),
            Col("Repeated-CV interval", 14), Col("ΔR² Tier 1 beyond Wave 1 outcome", 12, R3),
            Col("TabPFN R², full model", 10, R3)]
    spanners = [[("", 2), ("XGBoost R²", 3), ("", 4)]]
    return SITable(
        "S18", "table_s18_prediction_robustness.xlsx",
        "Robustness of the Tier 2 Increments: Repeated Cross-Validation, Specifications, Cohorts, "
        "Model Class, and Twin Pairs Held Out",
        [Block(cols, rows, spanners)],
        note=(
            "Out-of-fold R² under 5-fold cross-validation, with outcomes residualized on age within "
            "fold. Baseline model: Tier 1 for status, or the Wave 1 outcome plus Tier 1 for 9-year "
            "change. ΔR² Tier 2 is the full model minus the baseline model; ΔR² Tier 1 beyond the Wave 1 "
            "outcome is the baseline model minus the Wave 1 outcome alone. The repeated-CV interval "
            "gives the 2.5th and 97.5th percentiles of ΔR² Tier 2 over 15 repetitions with different "
            "fold assignments; it describes stability to the fold partition, not sampling error. "
            "Cohort rows fit the change model separately in MIDUS 2 (followed 2004-2009 to 2013-2014) "
            "and the MIDUS Refresher (2011-2014 to 2020-2022), each a single cross-validation run in "
            "about half the sample. TabPFN R² is from the same 177 predictors and folds. Frailty and "
            "cognition models always exclude complete twin pairs; the last rows refit the "
            "allostatic-load models with them removed from training, against the primary rows above, "
            f"and the two cross-sectional rankings shared {int(cs_held['top20_overlap'])} of their top 20 "
            f"predictors, with {LABEL[cs_held['top_feature_heldout']]} first in both. PhenoAge is the "
            "epigenetic benchmark (age-residualized PhenoAge). CV = cross-validation."
        ),
    )


# ── Assembly ───────────────────────────────────────────────────────────────
def build_all() -> list[SITable]:
    full = family_results()
    tables = [build_s1(), build_s2(), build_s3(), build_s4(full)]
    for key, layer, domain in (("S5", "cs", "Frailty"), ("S6", "cs", "Cognition"),
                               ("S7", "cs", "Allostatic load"), ("S8", "change", "Frailty"),
                               ("S9", "change", "Cognition"), ("S10", "change", "Allostatic load")):
        tables.append(build_family_table(key, layer, domain, full))
    tables += [build_s11(), build_s12(), build_s13(), build_s14(), build_s15(), build_s16(),
               build_s17(), build_s18()]
    return tables


def write_contents(ws, tables: list[SITable]) -> None:
    rows = [[t.key.replace("S", "Table S"), t.title, t.filename] for t in tables]
    contents = SITable(
        "S1-S18", "", "Supplementary Tables", [Block(
            [Col("Table", 12, left=True), Col("Title", 100, left=True), Col("Single-table file", 42, left=True)],
            rows)],
        note="Each table is on its own sheet of this workbook, named by its number, and in its own file.")
    write_table(ws, contents)
    ws["A1"] = "Supplementary Tables S1-S18"


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    tables = build_all()
    combined = Workbook()
    write_contents(combined.active, tables)
    combined.active.title = "Contents"
    for t in tables:
        wb = Workbook()
        wb.active.title = f"Table {t.key}"
        write_table(wb.active, t)
        wb.save(OUT / t.filename)
        print(f"Wrote {(OUT / t.filename).relative_to(ROOT)}")
        write_table(combined.create_sheet(t.key), t)
    path = OUT / "supplementary_tables_combined.xlsx"
    combined.save(path)
    print(f"Wrote {path.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
