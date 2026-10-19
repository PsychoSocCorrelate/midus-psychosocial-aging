#!/usr/bin/env python3
"""Figure 1: place the panels (A-G) on one PowerPoint slide at print size.

The panels are drawn at the size they print (fig1_subpanels.py), so
each is placed at 100% and never rescaled; a rescaled panel no longer has
6-8 pt type. The slide is the PNAS full width, 7.0 in. A PNG preview of the
same layout is written beside the deck. Each panel's sha256 is recorded in
fig1_panel_manifest.csv; --check reports panels regenerated since the deck
was built.

Layout, top to bottom:
  design   A cohorts (2.85 in) | B predictors (4.05 in)       neutral
           C outcomes (full width): the colour key for D-G
  results  the question D-G answer (full width, no letter)
           D cross-sectional | E longitudinal   (2.13 in tall, one R2 axis)
           F cross-domain    | G within-family  (1.66 in tall)

    uv run --with python-pptx python analysis/figures/fig1_deck.py
    uv run --with python-pptx python analysis/figures/fig1_deck.py --check
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))

FIG_DIR = ROOT / "manuscript" / "figures" / "fig1"
PANEL_DIR = FIG_DIR / "subpanels"
DECK = FIG_DIR / "fig1_panels.pptx"
MANIFEST = FIG_DIR / "fig1_panel_manifest.csv"
PREVIEW = FIG_DIR / "fig1_panels.png"

SLIDE_W = Inches(7.0)
PANEL_DPI = 600          # the panels' own resolution, so size = pixels / 600
ROW_GAP = 0.10           # inches between rows
SECTION_GAP = 0.22       # between the design rows and the result rows
LETTER_PT = 10
INK = RGBColor(0x22, 0x22, 0x22)
RULE = (0xD4, 0xD7, 0xDA)
LETTER_FONT = "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"

PANELS = {  # a one-letter key is a lettered panel; any other key is placed unlettered
    "A": "fig1_cohorts.png",
    "B": "fig1_predictors.png",
    "C": "fig1_outcomes.png",
    "question": "fig1_question.png",
    "D": "fig1_concurrent.png",
    "E": "fig1_change.png",
    "F": "fig1_crossdomain.png",
    "G": "fig1_cotwin.png",
}
ROWS = [  # (keys left to right, gap before the row)
    (("A", "B"), 0.0),
    (("C",), ROW_GAP),
    (("question",), SECTION_GAP),
    (("D", "E"), 0.04),
    (("F", "G"), ROW_GAP),
]


def lettered(key: str) -> bool:
    return len(key) == 1


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def size_in(path: Path) -> tuple[float, float]:
    """A panel's print size, inches, from its pixels at PANEL_DPI."""
    with Image.open(path) as im:
        return im.width / PANEL_DPI, im.height / PANEL_DPI


def layout() -> tuple[list[tuple[str, Path, float, float, float, float]], list[float], float]:
    """Every panel's (letter, path, x, y, w, h) in inches, the y of the rule
    between design and results, and the slide height. Rows are justified to
    the 7.0 in width: the first panel flush left, the last flush right."""
    placed, rules, y = [], [], 0.0
    W = SLIDE_W.inches
    for letters, gap in ROWS:
        if gap == SECTION_GAP:
            rules.append(y + gap / 2)
        y += gap
        sizes = [size_in(PANEL_DIR / PANELS[k]) for k in letters]
        if len(letters) == 1:
            xs = [0.0]
        else:
            xs = [0.0, W - sizes[1][0]]
        for k, x, (w, h) in zip(letters, xs, sizes):
            placed.append((k, PANEL_DIR / PANELS[k], x, y, w, h))
        y += max(h for _, h in sizes)
    return placed, rules, y


def render_png(placed, rules, height, dpi: int = 600) -> None:
    """The composite as a PNG at print resolution: the panels pasted at 100%,
    the letters set at LETTER_PT, as the deck places them."""
    W, H = round(SLIDE_W.inches * dpi), round(height * dpi)
    canvas = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(canvas)
    font = ImageFont.truetype(LETTER_FONT, round(LETTER_PT * dpi / 72))
    for y in rules:
        draw.line([(0, round(y * dpi)), (W, round(y * dpi))], fill=RULE,
                  width=max(1, round(0.5 * dpi / 72)))
    for letter, path, x, y, w, h in placed:
        with Image.open(path) as im:
            im = im.convert("RGB")
            if dpi != PANEL_DPI:
                im = im.resize((round(w * dpi), round(h * dpi)), Image.LANCZOS)
            canvas.paste(im, (round(x * dpi), round(y * dpi)))
        if lettered(letter):
            draw.text((round(x * dpi), round(y * dpi)), letter, font=font, fill=(0x22,) * 3)
    canvas.save(PREVIEW, dpi=(dpi, dpi))


def build() -> int:
    missing = [n for n in PANELS.values() if not (PANEL_DIR / n).exists()]
    if missing:
        for n in missing:
            print(f"  MISSING PANEL  {(PANEL_DIR / n).relative_to(ROOT)}")
        print("run fig1_subpanels first")
        return 2

    placed, rules, height = layout()

    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, Inches(height)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    for y in rules:
        line = slide.shapes.add_connector(1, Inches(0), Inches(y), SLIDE_W, Inches(y))
        line.line.color.rgb = RGBColor(*RULE)
        line.line.width = Pt(0.5)
    for letter, path, x, y, w, h in placed:
        slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
        if not lettered(letter):
            continue
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(0.17), Inches(0.18))
        tf = box.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        run = tf.paragraphs[0].add_run()
        run.text = letter
        run.font.size, run.font.bold, run.font.name = Pt(LETTER_PT), True, "Arial"
        run.font.color.rgb = INK

    DECK.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DECK)
    render_png(placed, rules, height)

    with open(MANIFEST, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["panel", "source", "sha256"])
        for letter, name in sorted(PANELS.items()):
            w.writerow([letter, f"subpanels/{name}", sha256(PANEL_DIR / name)])

    print(f"Saved: {DECK.relative_to(ROOT)} ({len(PANELS)} panels at "
          f"{SLIDE_W.inches:.1f} x {height:.2f} in)")
    for letter, path, x, y, w, h in placed:
        print(f"  {letter}  {path.name:24s} x {x:.2f}  y {y:.2f}  {w:.2f} x {h:.2f} in")
    print(f"Saved: {PREVIEW.relative_to(ROOT)}")
    print(f"Saved: {MANIFEST.relative_to(ROOT)}")
    return 0


def check() -> int:
    """Does the deck show the panels that are on disk now?

    Compare each placed panel's sha256 with the file on disk.
    """
    if not DECK.exists():
        print(f"  ABSENT   {DECK.relative_to(ROOT)}")
        return 1
    if not MANIFEST.exists():
        print(f"  NO MANIFEST  {MANIFEST.relative_to(ROOT)}: the deck cannot be "
              f"shown current without one")
        return 1

    stale, absent = [], []
    with open(MANIFEST) as fh:
        rows = list(csv.DictReader(fh))
    for row in rows:
        src = FIG_DIR / row["source"]
        if not src.exists():
            absent.append(row["source"])
        elif sha256(src) != row["sha256"]:
            stale.append(row["source"])

    for s in absent:
        print(f"  SOURCE GONE  {s}")
    for s in stale:
        print(f"  STALE        {s} has been regenerated since the deck was built")
    print(f"{len(rows)} panel(s): {len(rows) - len(stale) - len(absent)} current, "
          f"{len(stale)} stale, {len(absent)} missing")
    return 1 if (stale or absent) else 0


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--check", action="store_true",
                    help="report whether the deck shows the current panels, change nothing")
    args = ap.parse_args()
    return check() if args.check else build()


if __name__ == "__main__":
    sys.exit(main())
